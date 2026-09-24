"""
Filter specification for a graded stack.

Takes the lines Plot_Stack selected for a top-ranked stack (the CSV written by
Plot_Stack(..., selFile=...)) and works out the numbers needed to specify the
optics that separate and isolate them:

  Dichroics - one per adjacent pair of lines in a sub-stack, sorted by
    wavelength. The dichroic edge must lie between the RED edge of the lower
    line's Zeeman envelope and the BLUE edge of the upper line's envelope.
    e.g. Al III 360.193 | Mo I 386.410: the window runs from the top of the
    Al III envelope to the bottom of the Mo I envelope.

  Band-pass filters - one per selected line. The filter has to pass the line's
    own Zeeman envelope and block the closest radiation from any other modelled
    line on either side. The gap between the two is the room available for the
    filter's edges (after any CWL tolerance and angle-of-incidence shift).

Envelopes come from the same catalog Grade_Stack scores against
(Merge_Zeeman_Catalog: every measured h5 peak plus estimated envelopes for the
spreadsheet lines Curt has not computed), so the specification and the grading
cannot disagree about where a line's radiation sits.

    python FilterSpec.py TestStack_..._selected.csv
    python FilterSpec.py sel.csv --rank 2 --stack 2 3 --cwl-tol 0.1 --aoi 3 --save

or from python:

    from FilterSpec import Filter_Spec
    bp, dc = Filter_Spec('TestStack_..._selected.csv')

Outputs a formatted <prefix>_filter_spec.xlsx (Bandpass, Dichroics, Selection
and Notes sheets), a console summary, and one figure per sub-stack.

Limits of what this can tell you (also listed on the Notes sheet):
- Envelopes are the 1%-of-peak extent of each Zeeman pattern (LowHighCent).
  Relative line BRIGHTNESS is not modelled, so a faint neighbour limits the
  filter exactly as much as a bright one.
- Only modelled lines count as neighbours: lines in the spreadsheet or in
  Curt's h5 files. Molecular bands, continuum and unlisted impurity lines do
  not.
- Blocking far from the line (the rest of the detector's range) is a separate
  requirement from the nearest-neighbour limits reported here.
"""
import os
import io
import argparse
import contextlib

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.lines import Line2D
from matplotlib.gridspec import GridSpec

from StackMaker import (Merge_Zeeman_Catalog, Line_Envelopes, Load_H5_Spectra,
                        Plot_Zeeman_Spectra, Plot_Zeeman_Estimates, wave_to_rgb)

#same dichroic -> beamsplitter crossover (nm between line centres) Grade_Wvl uses
GBS = 2.0

#colours for the specification overlays. The spectra keep StackMaker's
#cycling colours and the lines their visible colour, as in SpeciesLines.
C_DICH = '#1f5aa6'    #dichroic edge window
C_BS = '#c26a00'      #window in the beamsplitter regime, or infeasible
C_FREE = '#2e7d32'    #free window between the nearest neighbours
C_NB = '#8b1a1a'      #nearest-neighbour radiation edge
C_INK = '0.15'

_ROMAN = ['', 'I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X',
          'XI', 'XII']


#---------------------------------------------------------------------------
# helpers
#---------------------------------------------------------------------------
def Ion_Label(el, ch):
    """'Al', 3 -> 'Al III'. The spreadsheet and the h5 files both use the
    spectroscopic number (He I = 1), so no offset is needed."""
    try:
        ch = int(ch)
    except (TypeError, ValueError):
        return f'{el} {ch}'
    return f'{el} {_ROMAN[ch]}' if 0 < ch < len(_ROMAN) else f'{el} {ch}'


def Catalog_Labels(files, cen, isEst):
    """
    Readable label for every catalog entry.
    'Al_3_2D_to_2Podd_360.193nm_term' + peak centre -> 'Al III 360.187'
    'Ni_1_471.442_EST_analytic-Lande'               -> 'Ni I 471.442 est'
    Measured entries carry the PEAK centre (what the spectrum shows), which can
    differ from the NIST wavelength by a few pm.
    """
    out = []
    for f, c, e in zip(files, cen, isEst):
        p = str(f).split('_')
        name = Ion_Label(p[0], p[1]) if len(p) > 1 else str(f)
        out.append(f'{name} {c:.3f}' + (' est' if e else ''))
    return out


def AOI_Shift(wvl, aoiDeg, nEff = 2.0):
    """
    Blue shift (nm, >= 0) of a thin-film interference filter's passband when
    light arrives at aoiDeg from normal:
        lambda(theta) = lambda0 * sqrt(1 - (sin(theta)/nEff)^2)
    H. A. Macleod, "Thin-Film Optical Filters", 4th ed., CRC Press (2010),
    chapter on band-pass filters at oblique incidence. nEff depends on the
    coating design (roughly 1.5-2.1 for common designs), so ask the vendor.
    The shift is always toward the blue, so it only eats into one side.
    """
    s = np.sin(np.radians(aoiDeg))/nEff
    return np.asarray(wvl, dtype=float)*(1.0 - np.sqrt(1.0 - s**2))


def Load_Selection(selFile, rank = 1):
    """
    Read the CSV Plot_Stack writes with selFile=... and keep one rank.
    Outputs a DataFrame sorted by stack, then wavelength.
    """
    df = pd.read_csv(selFile)
    ranks = sorted(df['rank'].unique())
    if rank not in ranks:
        raise ValueError(f'rank {rank} not in {selFile}; available: {ranks}')
    sel = df[df['rank'] == rank].copy()
    sel['note'] = sel['note'].fillna('').astype(str) if 'note' in sel else ''
    bad = sel['note'].str.startswith('NO MATCH')
    if bad.any():
        print(f'WARNING: {int(bad.sum())} selected line(s) could not be matched '
              f'to the line list when the selection was saved; their '
              f'wavelengths are float16 approximations:')
        print(sel.loc[bad, ['stack', 'species', 'wavelength_nm']].to_string(index=False))
    return sel.sort_values(['stack', 'wavelength_nm']).reset_index(drop=True)


def Load_Catalog(zFolder, xlsxFile):
    """The Zeeman catalog Grade_Stack scores against, plus readable labels.
    Nothing is written to zFolder."""
    cen, low, high, isEst, files = Merge_Zeeman_Catalog(
        zFolder, xlsxFile, saveCSV = False, verbose = False, returnLabels = True)
    cen, low, high = (np.asarray(x, dtype=float) for x in (cen, low, high))
    isEst = np.asarray(isEst, dtype=bool)
    return dict(cen=cen, low=low, high=high, isEst=isEst,
                file=np.asarray(files, dtype=object),
                lbl=np.asarray(Catalog_Labels(files, cen, isEst), dtype=object))


#---------------------------------------------------------------------------
# specifications
#---------------------------------------------------------------------------
def Bandpass_Spec(sel, cat, cwlTol = 0.0, aoiDeg = 0.0, nEff = 2.0,
                  tight = 0.1):
    """
    Band-pass requirement for every selected line.

    For each line:
      own envelope    [env_low, env_high] - the catalog entry Line_Envelopes
                      assigns to the line (the same rule Grade_Zeeman uses)
      blends          other envelopes that INTERSECT the own envelope. No
                      band-pass can separate these, so they are listed rather
                      than used as limits.
      blue neighbour  the highest red edge among the remaining envelopes lying
                      wholly below env_low - the closest radiation on the blue
                      side. The red neighbour is the mirror image.

    Turning that into a filter, allowing for a CWL tolerance of +-cwlTol and a
    cone of rays up to aoiDeg (which shifts the passband blue by dA):
      pass band    [env_low - cwlTol, env_high + cwlTol + dA]
                   what the filter must transmit at normal incidence so the
                   whole envelope passes for any tolerance or angle
      block below  blue_edge + cwlTol + dA   the nominal filter must already
      block above  red_edge  - cwlTol        be blocking at and beyond these
      edge room    distance from each pass-band edge to its blocking point,
                   i.e. the transition width that filter edge has to fit in.
                   The smaller of the two is what drives filter cost.
    With cwlTol = aoiDeg = 0 these reduce to the raw envelope and neighbour
    edges: the physics limit before any engineering margin.

    Inputs:
    - sel: DataFrame from Load_Selection
    - cat: dict from Load_Catalog
    - cwlTol: float - manufacturing CWL tolerance, +- nm
    - aoiDeg: float - maximum angle of incidence on the band-pass, degrees
    - nEff: float - filter effective index, for AOI_Shift
    - tight: float - edge room (nm) below which a line is flagged TIGHT
    Outputs:
    - DataFrame, one row per selected line
    """
    w = sel['wavelength_nm'].to_numpy(float)
    eLo, eHi, own = Line_Envelopes(w, cat['cen'], cat['low'], cat['high'],
                                   returnIdx = True)
    nCat = len(cat['cen'])
    rows = []
    for k, r in enumerate(sel.itertuples(index=False)):
        j = int(own[k])
        noEnv = j < 0
        lo, hi = (w[k], w[k]) if noEnv else (eLo[k], eHi[k])
        if noEnv:
            src = 'none in catalog'
        elif cat['isEst'][j]:
            src = 'estimated (' + str(cat['file'][j]).split('_EST_')[-1] + ')'
        else:
            src = 'measured (' + str(cat['file'][j]) + ')'

        other = np.ones(nCat, dtype=bool)
        if not noEnv:
            other[j] = False
        ov = other & (cat['low'] <= hi) & (cat['high'] >= lo)
        blue = other & ~ov & (cat['high'] < lo)
        red = other & ~ov & (cat['low'] > hi)

        if blue.any():
            jb = np.flatnonzero(blue)[np.argmax(cat['high'][blue])]
            bEdge, bLbl = cat['high'][jb], cat['lbl'][jb]
        else:
            bEdge, bLbl = np.nan, '(none modelled)'
        if red.any():
            jr = np.flatnonzero(red)[np.argmin(cat['low'][red])]
            rEdge, rLbl = cat['low'][jr], cat['lbl'][jr]
        else:
            rEdge, rLbl = np.nan, '(none modelled)'

        dA = float(AOI_Shift(w[k], aoiDeg, nEff))
        pLo, pHi = lo - cwlTol, hi + cwlTol + dA
        bBlue, bRed = bEdge + cwlTol + dA, rEdge - cwlTol
        roomB, roomR = pLo - bBlue, bRed - pHi
        rooms = [x for x in (roomB, roomR) if np.isfinite(x)]
        minRoom = min(rooms) if rooms else np.nan

        blends = list(cat['lbl'][ov])
        if noEnv:
            status = 'NO ENVELOPE: line not in the Zeeman catalog, width unknown'
        elif blends:
            status = (f'BLENDED with {len(blends)}: overlapping radiation '
                      f'cannot be rejected by a band-pass')
        elif rooms and minRoom <= 0:
            status = 'INFEASIBLE: tolerance/AOI exceed the gap to a neighbour'
        elif rooms and minRoom < tight:
            status = f'TIGHT: < {tight:g} nm for a filter edge'
        else:
            status = 'OK'

        rows.append(dict(
            stack=int(r.stack), line_label=f'{Ion_Label(r.species, r.ion)} {w[k]:.3f}',
            species=r.species, ion=r.ion, line_nm=w[k], prev_tok=bool(r.prev_tok),
            env_source=src, env_low_nm=lo, env_high_nm=hi, env_width_nm=hi - lo,
            blend_count=len(blends),
            blends='; '.join(blends[:6]) + (' ...' if len(blends) > 6 else ''),
            blue_neighbour=bLbl, blue_edge_nm=bEdge, blue_gap_nm=lo - bEdge,
            red_neighbour=rLbl, red_edge_nm=rEdge, red_gap_nm=rEdge - hi,
            cwl_nm=(pLo + pHi)/2, pass_low_nm=pLo, pass_high_nm=pHi,
            pass_width_nm=pHi - pLo, block_below_nm=bBlue, block_above_nm=bRed,
            edge_room_blue_nm=roomB, edge_room_red_nm=roomR,
            aoi_shift_nm=dA, status=status))
    return pd.DataFrame(rows)


def Dichroic_Spec(bp, chain = 'blue-first', gBS = GBS):
    """
    Dichroic requirement for every adjacent pair of lines in each sub-stack.

    The edge (50% point) has to fall inside
        [red edge of the lower line's envelope, blue edge of the upper line's]
    and everything the dichroic adds - its transition width, s/p polarisation
    splitting at 45 deg, cone-angle smear, manufacturing tolerance - has to fit
    in that window too. edge_50pct_nm is the window centre, which splits the
    budget evenly between the two lines.

    chain sets the topology. That decides what each dichroic has to reflect
    and transmit; the edge window itself is the same either way.
    - 'blue-first': long-pass dichroics, each reflects the bluest remaining
      line and transmits everything redder. D1 is the bluest split.
    - 'red-first': short-pass dichroics, each reflects the reddest remaining
      line and transmits everything bluer. D1 is the reddest split.

    Inputs:
    - bp: DataFrame from Bandpass_Spec (supplies each line's envelope)
    - chain: 'blue-first' or 'red-first'
    - gBS: float - line-centre gap (nm) below which Grade_Wvl assumes a
      beamsplitter is needed instead of a dichroic
    Outputs:
    - DataFrame, one row per dichroic
    """
    if chain not in ('blue-first', 'red-first'):
        raise ValueError("chain must be 'blue-first' or 'red-first'")
    rows = []
    for st, g in bp.groupby('stack', sort=True):
        g = g.sort_values('line_nm').reset_index(drop=True)
        n = len(g)
        for k in range(n - 1):
            L, U = g.iloc[k], g.iloc[k+1]
            eLo, eHi = L.env_high_nm, U.env_low_nm
            win, gap = eHi - eLo, U.line_nm - L.line_nm
            if chain == 'blue-first':
                name, order = f'D{k+1}', k + 1
                typ = 'long-pass: reflect lower, transmit upper'
                rLo, rHi = L.env_low_nm, L.env_high_nm
                tLo, tHi = U.env_low_nm, g.env_high_nm.iloc[k+1:].max()
            else:
                name, order = f'D{n-1-k}', n - 1 - k
                typ = 'short-pass: reflect upper, transmit lower'
                rLo, rHi = U.env_low_nm, U.env_high_nm
                tLo, tHi = g.env_low_nm.iloc[:k+1].min(), L.env_high_nm
            if win <= 0:
                status = 'INFEASIBLE: the two envelopes overlap'
            elif gap < gBS:
                status = f'BEAMSPLITTER regime: lines {gap:.2f} nm apart (< {gBS:g})'
            else:
                status = 'OK'
            rows.append(dict(
                stack=int(st), dichroic=name, _order=order,
                lower_line=L.line_label, upper_line=U.line_label,
                lower_env_high_nm=eLo, upper_env_low_nm=eHi,
                edge_window_nm=win, edge_50pct_nm=(eLo + eHi)/2,
                line_gap_nm=gap, type=typ,
                reflect_low_nm=rLo, reflect_high_nm=rHi,
                transmit_low_nm=tLo, transmit_high_nm=tHi, status=status))
    cols = ['stack', 'dichroic', 'lower_line', 'upper_line', 'lower_env_high_nm',
            'upper_env_low_nm', 'edge_window_nm', 'edge_50pct_nm', 'line_gap_nm',
            'type', 'reflect_low_nm', 'reflect_high_nm', 'transmit_low_nm',
            'transmit_high_nm', 'status']
    if not rows:
        return pd.DataFrame(columns=cols)
    return (pd.DataFrame(rows).sort_values(['stack', '_order'])[cols]
            .reset_index(drop=True))


#---------------------------------------------------------------------------
# tables
#---------------------------------------------------------------------------
def Print_Spec(bp, dc):
    """Short console summary, per sub-stack."""
    for st in sorted(bp['stack'].unique()):
        names = ', '.join(bp[bp['stack'] == st].line_label)
        print(f'\n=== stack {st}: {names} ===')
        d = dc[dc['stack'] == st]
        if len(d):
            print('  dichroics (edge must lie inside the window):')
        for r in d.itertuples():
            print(f'    {r.dichroic}: {r.lower_line:<17s}| {r.upper_line:<17s} '
                  f'window {r.lower_env_high_nm:8.3f} - {r.upper_env_low_nm:8.3f} nm '
                  f'({r.edge_window_nm:6.2f})  50% edge ~{r.edge_50pct_nm:8.2f}  '
                  f'{r.status}')
        print('  band-pass filters:')
        for r in bp[bp['stack'] == st].itertuples():
            print(f'    {r.line_label:<17s} envelope {r.env_low_nm:8.3f} - '
                  f'{r.env_high_nm:8.3f} ({r.env_width_nm:.3f} nm) | '
                  f'blue gap {r.blue_gap_nm:6.3f} to {r.blue_neighbour:<22s}| '
                  f'red gap {r.red_gap_nm:6.3f} to {r.red_neighbour:<22s}| '
                  f'{r.status}')


_NOTES = [
    ('Bandpass: env_low/high_nm', "The line's own Zeeman envelope (1% of peak) from the catalog Grade_Stack uses. env_source says whether it is a measured h5 peak or an estimate."),
    ('Bandpass: blends', "Other envelopes that intersect the line's own. No band-pass can reject these, so they are listed rather than used as limits."),
    ('Bandpass: blue/red_neighbour, *_edge_nm', 'Closest modelled radiation on each side: the nearest envelope edge lying wholly outside the own envelope.'),
    ('Bandpass: *_gap_nm', 'Raw distance from the own envelope to that neighbour edge, before any tolerance.'),
    ('Bandpass: pass_low/high_nm, cwl_nm, pass_width_nm', 'Region the filter must transmit at normal incidence: the envelope widened by cwlTol on the blue and cwlTol + AOI shift on the red. cwl_nm is its centre and pass_width_nm the minimum flat-top width.'),
    ('Bandpass: block_below/above_nm', 'The nominal filter must already be blocking at and beyond these, so that a +-cwlTol shift and the AOI blue shift still block the neighbours.'),
    ('Bandpass: edge_room_*_nm', 'Space for each filter edge to go from pass to block. The smaller one drives filter steepness and cost; <= 0 means the tolerances do not fit.'),
    ('Dichroics: edge window', 'From the red edge of the lower line envelope to the blue edge of the upper line envelope. The 50% edge, transition width, s/p splitting at 45 deg, cone smear and CWL tolerance must all fit inside it.'),
    ('Dichroics: edge_50pct_nm', 'Window centre, splitting the budget evenly between the two lines.'),
    ('Dichroics: reflect/transmit band', 'What this dichroic must reflect and transmit for the chosen chain topology (see parameters above).'),
    ('Not modelled', 'Relative line brightness (a faint neighbour limits as much as a bright one), unlisted impurity lines, molecular bands and continuum, and blocking across the full detector range.'),
]


def _format_sheet(ws, df):
    """Bold wrapped header, frozen top row, sensible widths, 3 dp numbers, and
    a highlight on any status that is not OK."""
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    flag = PatternFill('solid', fgColor='FFF3CD')
    for i, col in enumerate(df.columns, 1):
        L = get_column_letter(i)
        vals = [f'{v:.3f}' if isinstance(v, float) else str(v) for v in df[col]]
        ws.column_dimensions[L].width = min(max([len(str(col))*0.8] +
                                                [len(v) for v in vals]) + 2, 70)
        if df[col].dtype.kind == 'f':
            for cell in ws[L][1:]:
                cell.number_format = '0.000'
        if col == 'status':
            for cell in ws[L][1:]:
                if not str(cell.value).startswith('OK'):
                    cell.fill = flag


def Write_Spec_Xlsx(bp, dc, sel, params, path):
    """Formatted workbook with Bandpass, Dichroics, Selection and Notes sheets."""
    notes = pd.DataFrame([(k, str(v)) for k, v in params.items()] + [('', '')] +
                         _NOTES, columns=['item', 'value / definition'])
    sheets = [('Bandpass', bp), ('Dichroics', dc), ('Selection', sel),
              ('Notes', notes)]
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
            _format_sheet(xw.sheets[name], df)
        xw.sheets['Notes'].column_dimensions['B'].width = 110
    print(f'wrote {path}')


#---------------------------------------------------------------------------
# plot
#---------------------------------------------------------------------------
def Plot_Spec(bp, dc, cat, zFolder, title = '', stacks = None, window = 1.0,
              maxWindow = 6.0, specH = 3.0, cmap = 'hsv', savePrefix = None,
              show = True):
    """
    One figure per sub-stack, laid out like SpeciesLines.Plot_Species:

    - overview strip: the selected lines in their visible colour, each
      dichroic's edge window shaded (blue = dichroic, amber = beamsplitter
      regime or infeasible) with its suggested 50% edge dashed
    - one zoom panel per line: the measured Zeeman spectra and estimated
      envelopes nearby, the free window between the nearest neighbours
      (green), the neighbour edges (dark red), and the limiting filter (black
      dashed trapezoid: flat over the pass band, fully blocking at
      block_below/above_nm). Its sloped sides are the SHALLOWEST edges that
      still meet the spec - a real filter's edges must be at least this steep.
      Transmission and the spectra share the same normalised 0-1 scale.

    Inputs:
    - bp, dc: DataFrames from Bandpass_Spec / Dichroic_Spec
    - cat: dict from Load_Catalog (for the estimated envelopes)
    - zFolder: str - h5 folder for the measured spectra
    - title: str - figure title prefix
    - stacks: list of int or None (all)
    - window: float - minimum half-width (nm) of a zoom panel
    - maxWindow: float - maximum half-width. A neighbour further away than
      this gets an arrow at the panel edge instead of widening the panel.
    - specH, cmap: as SpeciesLines
    - savePrefix: str or None - write <savePrefix>_stack<N>.png
    - show: bool - call plt.show()
    Outputs:
    - list of figures
    """
    with contextlib.redirect_stdout(io.StringIO()):   #LowHighCent is chatty
        spectra = Load_H5_Spectra(zFolder)
    est = cat['isEst']
    eCen, eLow, eHigh = cat['cen'][est], cat['low'][est], cat['high'][est]
    eLbl = list(cat['lbl'][est])

    figs = []
    for st in (stacks or sorted(bp['stack'].unique())):
        g = bp[bp['stack'] == st].sort_values('line_nm').reset_index(drop=True)
        d = dc[dc['stack'] == st]
        n = len(g)
        if n == 0:
            continue
        fig = plt.figure(figsize=(4.4*max(n, 3), 8.2))
        gs = GridSpec(2, n, figure=fig, height_ratios=[1, 3.2], hspace=0.45,
                      wspace=0.20)

        # ---------------- overview: dichroic windows ----------------
        axO = fig.add_subplot(gs[0, :])
        for m, r in enumerate(d.itertuples()):
            col = C_DICH if r.status.startswith('OK') else C_BS
            #windows sit back to back (separated only by a line envelope), so
            #alternate the tone to keep neighbouring windows distinguishable
            axO.axvspan(r.lower_env_high_nm, r.upper_env_low_nm, color=col,
                        alpha=0.10 if m % 2 else 0.24, lw=0, zorder=0)
            axO.axvline(r.edge_50pct_nm, color=col, ls='--', lw=1.1, zorder=1)
            axO.text(r.edge_50pct_nm, 0.10 + 0.34*(m % 2),
                     f'{r.dichroic}\n{r.edge_50pct_nm:.2f}', ha='center',
                     va='bottom', fontsize=8, color=C_INK,
                     bbox=dict(boxstyle='round,pad=0.18', fc='white', ec=col,
                               lw=0.8, alpha=0.9), zorder=3)
        span = max(g.line_nm.max() - g.line_nm.min(), 1.0)
        row, prev = 0, -np.inf
        for r in g.itertuples():
            #second row when a label would collide with its left neighbour
            row = (1 - row) if r.line_nm - prev < 0.035*span else 0
            prev = r.line_nm
            axO.vlines(r.line_nm, 0, 1 + 0.13*row, color=wave_to_rgb(r.line_nm),
                       lw=2.2, zorder=2)
            axO.text(r.line_nm, 1.03 + 0.13*row, r.line_label.rsplit(' ', 1)[0],
                     ha='center', va='bottom', fontsize=8.5, color=C_INK,
                     fontweight='bold' if r.prev_tok else 'normal')
        lo, hi = g.env_low_nm.min(), g.env_high_nm.max()
        pad = max(0.04*(hi - lo), 4)
        axO.set_xlim(lo - pad, hi + pad)
        axO.set_ylim(0, 1.42)
        axO.set_yticks([])
        axO.set_xlabel('Wavelength (nm)', fontsize=9)
        axO.grid(axis='x', color='0.9', lw=0.5)
        axO.set_title(f'{title}stack {st}: dichroic edge windows '
                      f'(red edge of lower envelope to blue edge of upper)',
                      fontsize=11, fontweight='bold')

        # ---------------- zoom: band-pass per line ----------------
        for k, r in enumerate(g.itertuples()):
            ax = fig.add_subplot(gs[1, k])
            w = r.line_nm
            reach = np.nanmax([w - r.blue_edge_nm, r.red_edge_nm - w, window])
            half = float(min(reach*1.15, maxWindow))
            x0, x1 = w - half, w + half

            drawn = Plot_Zeeman_Spectra(ax, spectra, [w], window=half,
                                        height=specH, cmap=cmap)
            if len(eCen):
                Plot_Zeeman_Estimates(ax, eCen, eLow, eHigh, eLbl, [w],
                                      window=half, height=specH, cmap=cmap,
                                      cStart=len(drawn))

            #free window between the nearest neighbours
            fLo = r.blue_edge_nm if np.isfinite(r.blue_edge_nm) else x0
            fHi = r.red_edge_nm if np.isfinite(r.red_edge_nm) else x1
            ax.axvspan(max(fLo, x0), min(fHi, x1), color=C_FREE, alpha=0.09,
                       lw=0, zorder=0)

            #own envelope edges
            for e in (r.env_low_nm, r.env_high_nm):
                ax.axvline(e, color=C_INK, ls=':', lw=1.0, zorder=7)

            #nearest-neighbour edges: a line if on the panel, else an arrow
            #labels go INSIDE the free window (so they never run into the next
            #panel), the red one a row lower so the two cannot collide
            for edge, lbl, side, yN in ((r.blue_edge_nm, r.blue_neighbour, -1, specH*1.80),
                                        (r.red_edge_nm, r.red_neighbour, +1, specH*1.52)):
                if not np.isfinite(edge):
                    continue
                txt = f'{lbl}\n{edge:.3f} nm'
                if x0 <= edge <= x1:
                    ax.axvline(edge, color=C_NB, lw=1.6, zorder=8)
                    ax.text(edge - side*0.03*half, yN, txt, color=C_NB,
                            fontsize=7, ha='left' if side < 0 else 'right',
                            va='bottom', zorder=9,
                            bbox=dict(boxstyle='round,pad=0.12', fc='white',
                                      ec='none', alpha=0.75))
                else:
                    xe = x0 if side < 0 else x1
                    ax.annotate(txt, xy=(xe, yN), xytext=(xe - side*0.30*half, yN),
                                color=C_NB, fontsize=7, va='center',
                                ha='left' if side < 0 else 'right',
                                arrowprops=dict(arrowstyle='->', color=C_NB, lw=1.2),
                                zorder=9)

            #limiting filter: the shallowest edges that still meet the spec
            pts = [r.block_below_nm, r.pass_low_nm, r.pass_high_nm, r.block_above_nm]
            if (np.isfinite(pts).all() and r.edge_room_blue_nm > 0
                    and r.edge_room_red_nm > 0):
                ax.plot(pts, [0, specH, specH, 0], color='k', ls='--', lw=1.5,
                        zorder=10)
            ax.vlines(w, 0, specH*1.45, color=wave_to_rgb(w), lw=1.6, zorder=8)

            ax.set_xlim(x0, x1)
            ax.set_ylim(0, specH*2.75)
            ax.set_yticks([])
            ax.set_xlabel('Wavelength (nm)', fontsize=9)
            ax.grid(axis='x', color='0.9', lw=0.5)
            ax.set_title(f'{r.line_label} nm  [{r.status.split(":")[0]}]',
                         fontsize=9.5,
                         fontweight='bold' if r.prev_tok else 'normal')
            info = (f'envelope {r.env_width_nm:.3f} nm, {r.env_source.split(" ")[0]}\n'
                    f'CWL {r.cwl_nm:.3f}  pass {r.pass_low_nm:.3f}-{r.pass_high_nm:.3f}\n'
                    f'block <= {r.block_below_nm:.3f}  >= {r.block_above_nm:.3f}\n'
                    f'edge room blue {r.edge_room_blue_nm:.3f}  red {r.edge_room_red_nm:.3f}')
            if r.blend_count:
                info += f'\nblended with {r.blend_count}'
            ax.text(0.02, 0.985, info, transform=ax.transAxes, fontsize=6.8,
                    va='top', ha='left', color=C_INK, family='monospace',
                    bbox=dict(boxstyle='round,pad=0.3', fc='white', ec='0.7',
                              alpha=0.92), zorder=11)

        leg = [Patch(facecolor=C_DICH, alpha=0.3, label='dichroic edge window'),
               Patch(facecolor=C_BS, alpha=0.4, label='beamsplitter regime / infeasible'),
               Line2D([0], [0], color=C_INK, ls=':', lw=1, label='own Zeeman envelope'),
               Patch(facecolor=C_FREE, alpha=0.2, label='free window'),
               Line2D([0], [0], color=C_NB, lw=1.6, label='nearest modelled radiation'),
               Line2D([0], [0], color='k', ls='--', lw=1.5,
                      label='limiting filter: edges no shallower (T 0-1)'),
               Patch(facecolor='0.8', edgecolor='0.4', ls='--', hatch='///',
                     label='estimated envelope (no h5)')]
        fig.legend(handles=leg, ncol=4, fontsize=8.5, loc='lower center',
                   bbox_to_anchor=(0.5, -0.035), frameon=False)
        if savePrefix:
            p = f'{savePrefix}_stack{st}.png'
            fig.savefig(p, dpi=140, bbox_inches='tight')
            print('saved', p)
        figs.append(fig)
    if show:
        plt.show()
    return figs


#---------------------------------------------------------------------------
# driver
#---------------------------------------------------------------------------
def Filter_Spec(selFile, rank = 1, stacks = None, cwlTol = 0.0, aoiDeg = 0.0,
                nEff = 2.0, chain = 'blue-first', gBS = GBS, tight = 0.1,
                xlsxFile = None, zFolder = None, outPrefix = None, plot = True,
                savePlots = False, show = True):
    """
    Everything in one call: load a ranked selection, compute the dichroic and
    band-pass specs, print them, write the workbook, and plot.

    Inputs:
    - selFile: str - CSV from Plot_Stack(selFile=...)
    - rank: int - which ranked combination to specify (1 = best)
    - stacks: list of int or None - sub-stacks to include (None = all)
    - cwlTol, aoiDeg, nEff, tight: see Bandpass_Spec
    - chain, gBS: see Dichroic_Spec
    - xlsxFile, zFolder: override the line list / h5 folder recorded in the
      selection file (by default the ones the stack was plotted with)
    - outPrefix: str - output name prefix, default <selFile>_rank<N>
    - plot, savePlots, show: figure control
    Outputs:
    - bp, dc: the band-pass and dichroic DataFrames
    """
    sel = Load_Selection(selFile, rank)
    xlsxFile = xlsxFile or sel['xlsx'].iloc[0]
    zFolder = zFolder or sel['zfolder'].iloc[0]
    if stacks:
        sel = sel[sel['stack'].isin(stacks)].reset_index(drop=True)
    cat = Load_Catalog(zFolder, xlsxFile)

    bp = Bandpass_Spec(sel, cat, cwlTol, aoiDeg, nEff, tight)
    dc = Dichroic_Spec(bp, chain, gBS)

    outPrefix = outPrefix or os.path.splitext(selFile)[0] + f'_rank{rank}'
    r0 = sel.iloc[0]
    params = {'selection file': selFile, 'rank': rank,
              'combination': int(r0['combination']),
              'total score': r0['total_score'], 'score weights': r0['weights'],
              'line list (xlsx)': xlsxFile, 'Zeeman h5 folder': zFolder,
              'CWL tolerance (+- nm)': cwlTol,
              'max AOI on band-pass (deg)': aoiDeg, 'filter n_eff': nEff,
              'dichroic chain': chain, 'beamsplitter gap gBS (nm)': gBS,
              'tight edge-room flag (nm)': tight,
              'catalog entries': f'{int((~cat["isEst"]).sum())} measured + '
                                 f'{int(cat["isEst"].sum())} estimated'}

    print(f'rank {rank} (combination {params["combination"]}, total score '
          f'{params["total score"]}) | CWL tol +-{cwlTol} nm | AOI {aoiDeg} deg '
          f'| chain {chain}')
    Print_Spec(bp, dc)
    Write_Spec_Xlsx(bp, dc, sel, params, outPrefix + '_filter_spec.xlsx')
    if plot:
        Plot_Spec(bp, dc, cat, zFolder, title=f'rank {rank}, ',
                  savePrefix=outPrefix if savePlots else None, show=show)
    return bp, dc


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('selfile', help='CSV written by Plot_Stack(selFile=...)')
    ap.add_argument('--rank', type=int, default=1)
    ap.add_argument('--stack', type=int, nargs='*', help='sub-stack number(s)')
    ap.add_argument('--cwl-tol', type=float, default=0.0,
                    help='band-pass CWL tolerance, +- nm')
    ap.add_argument('--aoi', type=float, default=0.0,
                    help='max angle of incidence on the band-pass, degrees')
    ap.add_argument('--neff', type=float, default=2.0,
                    help='band-pass effective index for the AOI shift')
    ap.add_argument('--chain', default='blue-first',
                    choices=['blue-first', 'red-first'])
    ap.add_argument('--xlsx', help='override the line list in the selection file')
    ap.add_argument('--zfolder', help='override the h5 folder in the selection file')
    ap.add_argument('--out', help='output prefix')
    ap.add_argument('--save', action='store_true',
                    help='save the figures as PNG instead of showing them')
    a = ap.parse_args()
    Filter_Spec(a.selfile, rank=a.rank, stacks=a.stack, cwlTol=a.cwl_tol,
                aoiDeg=a.aoi, nEff=a.neff, chain=a.chain, xlsxFile=a.xlsx,
                zFolder=a.zfolder, outPrefix=a.out, savePlots=a.save,
                show=not a.save)
